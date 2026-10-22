"""
Laporan otomatis titik panas – Desa & Hutan Desa Nanga Lauk.

Membuat laporan utuh (Word .docx dan PDF) dari basis data titik panas:
  * Laporan lengkap 2001–sekarang
  * Laporan tahunan untuk setiap tahun
Isi: ringkasan, metodologi & sumber data, peta, grafik, tabel rekap, daftar detail
setiap deteksi, kejadian api, kamus data, dan jejak audit. Tanpa logo.

Dipanggil oleh scripts/pipeline.py: build_reports(recs, events, years, full=True)
"""
import os, io, json, glob, datetime, tempfile
from collections import Counter, defaultdict
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.lines import Line2D
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
                                PageBreak, KeepTogether)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(DATA, "laporan")
ZONES = ["Hutan Desa", "Desa Nanga Lauk", "Radius 5 km", "Radius 5-25 km", "Kapuas Hulu", "Luar Kapuas Hulu"]
ZCOL = {"Hutan Desa": "#C21D1D", "Desa Nanga Lauk": "#E2621B", "Radius 5 km": "#E8A317", "Radius 5-25 km": "#C9B21E",
        "Kapuas Hulu": "#8E7A9E", "Luar Kapuas Hulu": "#9AA49E"}
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
INK = "#1D2A22"; MUTED = "#5B6A61"; LINE = "#D3DBD5"; HEAD = "#22382B"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": MUTED, "axes.labelcolor": INK,
                     "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False})

geo = json.load(open(os.path.join(DATA, "geo.json")))

def fam(r):
    if r["instrument"] == "MODIS":
        return "MODIS"
    return "VIIRS S-NPP" if "S-NPP" in r["sensor"] else ("VIIRS NOAA-20" if "NOAA-20" in r["sensor"] else "VIIRS NOAA-21")

def valid(r):
    return r.get("type") not in ("1", "2", "3")

def num(v):
    try: return float(v)
    except Exception: return 0.0

def fmt(n):
    return f"{n:,}".replace(",", ".")

def conf_txt(r):
    if r["instrument"] == "VIIRS":
        return {"l": "rendah", "n": "nominal", "h": "tinggi"}.get(r["confidence"], r["confidence"])
    return f'{r["confidence"]}%'

# ------------------------------------------------------------------ gambar
def rings(geom):
    polys = [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]
    for p in polys:
        for r in p:
            yield r

def draw_fc(ax, fc, **kw):
    for f in fc["features"]:
        g = f["geometry"]
        if g["type"] in ("Polygon", "MultiPolygon"):
            for r in rings(g):
                xs, ys = zip(*r); ax.plot(xs, ys, **kw)
        elif g["type"] in ("LineString", "MultiLineString"):
            lines = [g["coordinates"]] if g["type"] == "LineString" else g["coordinates"]
            for l in lines:
                xs, ys = zip(*l); ax.plot(xs, ys, **kw)

def fig_map25(rows, path, title):
    fig, ax = plt.subplots(figsize=(7.2, 6.4))
    b = [c for f in geo["buf25"]["features"] for r in rings(f["geometry"]) for c in r]
    x0, x1 = min(c[0] for c in b) - 0.02, max(c[0] for c in b) + 0.02
    y0, y1 = min(c[1] for c in b) - 0.02, max(c[1] for c in b) + 0.02
    ax.set_facecolor("#F4F6F3")
    draw_fc(ax, geo["kab"], color="#9AA49E", lw=0.6, ls=(0, (3, 3)))
    draw_fc(ax, geo["river"], color="#4A90C8", lw=0.8, alpha=0.8)
    draw_fc(ax, geo["buf25"], color=ZCOL["Radius 5-25 km"], lw=1, ls=(0, (6, 4)))
    draw_fc(ax, geo["buf5"], color=ZCOL["Radius 5 km"], lw=1, ls=(0, (6, 3)))
    draw_fc(ax, geo["desa"], color="#222", lw=1.3)
    draw_fc(ax, geo["hd"], color="#C21D1D", lw=1.6)
    for z in reversed(ZONES[:4]):
        pts = [r for r in rows if r["zona"] == z]
        if pts:
            ax.scatter([num(r["longitude"]) for r in pts], [num(r["latitude"]) for r in pts], s=14 if z in ZONES[:2] else 8,
                       c=ZCOL[z], edgecolors="black", linewidths=0.3, zorder=5, label=f"{z} ({len(pts)})")
    ax.set_xlim(x0, x1); ax.set_ylim(y0, y1); ax.set_aspect(1 / 0.99985)
    ax.set_title(title, fontsize=10, color=INK, loc="left")
    h = [Line2D([], [], color="#C21D1D", lw=1.6, label="Hutan Desa"), Line2D([], [], color="#222", lw=1.3, label="Desa Nanga Lauk"),
         Line2D([], [], color=ZCOL["Radius 5 km"], ls="--", label="Radius 5 km"), Line2D([], [], color=ZCOL["Radius 5-25 km"], ls="--", label="Radius 25 km"),
         Line2D([], [], color="#4A90C8", label="Sungai utama")]
    l1 = ax.legend(handles=h, loc="lower left", fontsize=7, frameon=True); ax.add_artist(l1)
    if rows:
        ax.legend(loc="upper right", fontsize=7, title="Titik panas", title_fontsize=7, frameon=True)
    ax.tick_params(labelsize=7); ax.set_xlabel("Bujur (°BT)", fontsize=8); ax.set_ylabel("Lintang (°LU)", fontsize=8)
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig)

def fig_mapdesa(rows, path, title):
    fig, ax = plt.subplots(figsize=(7.2, 6.4))
    b = [c for f in geo["buf5"]["features"] for r in rings(f["geometry"]) for c in r]
    x0, x1 = min(c[0] for c in b), max(c[0] for c in b)
    y0, y1 = min(c[1] for c in b), max(c[1] for c in b)
    img = os.path.join(ROOT, "assets", "s2_20260810.jpg")
    if os.path.exists(img):
        (la0, lo0), (la1, lo1) = geo["img_bounds"]
        ax.imshow(plt.imread(img), extent=[lo0, lo1, la0, la1], zorder=0)
    draw_fc(ax, geo["grid"], color="white", lw=0.25, alpha=0.6)
    draw_fc(ax, geo["buf5"], color=ZCOL["Radius 5 km"], lw=1.1, ls=(0, (6, 3)))
    draw_fc(ax, geo["desa"], color="#FFE14D", lw=1.4)
    draw_fc(ax, geo["hd"], color="#FF3B30", lw=1.8)
    for z in reversed(ZONES[:3]):
        pts = [r for r in rows if r["zona"] == z]
        if pts:
            ax.scatter([num(r["longitude"]) for r in pts], [num(r["latitude"]) for r in pts], s=26 if z in ZONES[:2] else 16,
                       c=ZCOL[z], edgecolors="white", linewidths=0.6, zorder=5, label=f"{z} ({len(pts)})")
    ax.set_xlim(x0, x1); ax.set_ylim(y0, y1); ax.set_aspect(1 / 0.99985); ax.set_facecolor("#DDE3DE")
    ax.set_title(title, fontsize=10, color=INK, loc="left")
    if rows:
        ax.legend(loc="upper right", fontsize=7, title="Titik panas", title_fontsize=7, frameon=True)
    ax.tick_params(labelsize=7); ax.set_xlabel("Bujur (°BT)", fontsize=8); ax.set_ylabel("Lintang (°LU)", fontsize=8)
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig)

def fig_annual(rows, family, path, title):
    zones = ZONES[:4]
    years = sorted({r["acq_date"][:4] for r in rows if fam(r) == family})
    if not years:
        return False
    cnt = Counter((r["acq_date"][:4], r["zona"]) for r in rows if fam(r) == family)
    fig, ax = plt.subplots(figsize=(7.2, 3.0)); bottom = [0] * len(years)
    for z in zones:
        v = [cnt[(y, z)] for y in years]
        ax.bar(years, v, bottom=bottom, color=ZCOL[z], label=z, width=0.75)
        bottom = [a + b for a, b in zip(bottom, v)]
    for i, t in enumerate(bottom):
        if t: ax.text(i, t, str(t), ha="center", va="bottom", fontsize=6.5, color=INK)
    ax.set_title(title, fontsize=10, color=INK, loc="left"); ax.set_ylabel("Jumlah deteksi")
    ax.tick_params(axis="x", labelrotation=90, labelsize=7); ax.grid(axis="y", color=LINE, lw=0.6); ax.set_axisbelow(True)
    ax.legend(fontsize=7, ncol=4, frameon=False, loc="upper left")
    ax.set_ylim(0, max(bottom) * 1.18 + 1)
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig); return True

def fig_monthly(rows, path, title, year=None):
    zones = ZONES[:4]
    cnt = Counter((int(r["acq_date"][5:7]), r["zona"]) for r in rows)
    fig, ax = plt.subplots(figsize=(7.2, 2.8)); bottom = [0] * 12
    for z in zones:
        v = [cnt[(m, z)] for m in range(1, 13)]
        ax.bar(BULAN, v, bottom=bottom, color=ZCOL[z], label=z, width=0.7)
        bottom = [a + b for a, b in zip(bottom, v)]
    for i, t in enumerate(bottom):
        if t: ax.text(i, t, str(t), ha="center", va="bottom", fontsize=7, color=INK)
    ax.set_title(title, fontsize=10, color=INK, loc="left"); ax.set_ylabel("Jumlah deteksi")
    ax.grid(axis="y", color=LINE, lw=0.6); ax.set_axisbelow(True); ax.legend(fontsize=7, ncol=4, frameon=False, loc="upper left")
    ax.set_ylim(0, max(bottom + [1]) * 1.2 + 1)
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig)

def fig_context(rows, path):
    near = [r for r in rows if r["zona"] in ZONES[:3]]
    if not near:
        return False
    fig, axs = plt.subplots(1, 3, figsize=(7.4, 2.9))
    for ax, key, t in zip(axs, ["fungsi_kawasan", "tutupan_2026", "gambut"], ["Fungsi kawasan", "Tutupan lahan 2026", "Gambut"]):
        c = Counter((r[key] or "Tidak ada / di luar data").split(" (")[0] if key != "fungsi_kawasan" else (r[key] or "Di luar data") for r in near)
        items = c.most_common(6)[::-1]
        ax.barh([i[0][:28] for i in items], [i[1] for i in items], color="#7E9C87")
        for y, (k, v) in enumerate(items):
            ax.text(v, y, f" {v}", va="center", fontsize=7, color=INK)
        ax.set_title(t, fontsize=9, color=INK, loc="left"); ax.tick_params(labelsize=6.5); ax.set_xlim(0, max(i[1] for i in items) * 1.3)
    fig.suptitle("Konteks lokasi titik panas di desa & radius 5 km", fontsize=10, color=INK, x=0.01, ha="left")
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig); return True

def fig_hour(rows, path):
    near = [r for r in rows if r["zona"] in ZONES[:4]]
    if not near:
        return False
    c = Counter((int(r["waktu_wib"][11:13])) for r in near)
    fig, ax = plt.subplots(figsize=(7.2, 2.4))
    ax.bar(range(24), [c[h] for h in range(24)], color="#7E9C87", width=0.8)
    ax.set_xticks(range(0, 24, 2)); ax.set_xlabel("Jam deteksi (WIB)"); ax.set_ylabel("Jumlah")
    ax.set_title("Jam lintasan satelit saat titik panas terdeteksi (radius 25 km)", fontsize=10, color=INK, loc="left")
    ax.grid(axis="y", color=LINE, lw=0.6); ax.set_axisbelow(True)
    fig.tight_layout(); fig.savefig(path, dpi=140); plt.close(fig); return True

# ------------------------------------------------------------------ isi laporan (struktur netral)
def content(recs, events, year=None):
    rows = [r for r in recs.values() if valid(r) and (year is None or r["acq_date"][:4] == str(year))]
    focus = [r for r in rows if r["zona"] in ZONES[:4]]
    near = sorted([r for r in rows if r["zona"] in ZONES[:3]], key=lambda r: (r["acq_date"], r["acq_time"]))
    ring = sorted([r for r in rows if r["zona"] == "Radius 5-25 km"], key=lambda r: (r["acq_date"], r["acq_time"]))
    evs = [e for e in events if e["zona_terdekat"] in ZONES[:3] and (year is None or e["mulai_wib"][:4] == str(year))]
    zc = Counter(r["zona"] for r in rows)
    span = (min((r["acq_date"] for r in rows), default="-"), max((r["acq_date"] for r in rows), default="-"))
    tmp = tempfile.mkdtemp()
    F = {}
    lbl = f"tahun {year}" if year else f"{span[0][:4]}–{span[1][:4]}"
    F["map25"] = os.path.join(tmp, "map25.png"); fig_map25(focus, F["map25"], f"Titik panas radius 25 km dari Desa Nanga Lauk, {lbl}")
    F["mapdesa"] = os.path.join(tmp, "mapdesa.png"); fig_mapdesa(near, F["mapdesa"], f"Titik panas di desa & radius 5 km, {lbl} (latar Sentinel-2 10 Agu 2026)")
    if year is None:
        if fig_annual(focus, "MODIS", os.path.join(tmp, "am.png"), "Deteksi per tahun – MODIS 1 km (seri konsisten sejak 2001)"): F["annual_modis"] = os.path.join(tmp, "am.png")
        if fig_annual(focus, "VIIRS S-NPP", os.path.join(tmp, "av.png"), "Deteksi per tahun – VIIRS S-NPP 375 m (sejak 2012)"): F["annual_viirs"] = os.path.join(tmp, "av.png")
    F["monthly"] = os.path.join(tmp, "mo.png")
    fig_monthly(focus, F["monthly"], ("Sebaran bulanan seluruh tahun (pola musim)" if year is None else f"Deteksi per bulan, {year}") + " – radius 25 km, semua sensor")
    if fig_context(rows, os.path.join(tmp, "ctx.png")): F["context"] = os.path.join(tmp, "ctx.png")
    if fig_hour(rows, os.path.join(tmp, "hr.png")): F["hour"] = os.path.join(tmp, "hr.png")
    # tabel rekap zona per sensor
    fams = ["MODIS", "VIIRS S-NPP", "VIIRS NOAA-20", "VIIRS NOAA-21"]
    zt = [["Zona"] + fams + ["Semua"]]
    for z in ZONES[:5]:
        zt.append([z] + [fmt(sum(1 for r in rows if r["zona"] == z and fam(r) == f)) for f in fams] + [fmt(zc[z])])
    # rekap tahunan (laporan lengkap) atau bulanan (tahunan)
    if year is None:
        yrs = sorted({r["acq_date"][:4] for r in rows})
        tt = [["Tahun", "Hutan Desa", "Desa", "≤ 5 km", "5–25 km", "Kapuas Hulu*", "Sensor tersedia"]]
        for y in yrs:
            ry = [r for r in rows if r["acq_date"][:4] == y]; c = Counter(r["zona"] for r in ry)
            tt.append([y, str(c["Hutan Desa"]), str(c["Desa Nanga Lauk"]), str(c["Radius 5 km"]), str(c["Radius 5-25 km"]),
                       fmt(sum(c[z] for z in ZONES[:5])), ", ".join(sorted({fam(r) for r in ry}))])
    else:
        tt = [["Bulan", "Hutan Desa", "Desa", "≤ 5 km", "5–25 km", "Kapuas Hulu*"]]
        for m in range(1, 13):
            rm = [r for r in rows if int(r["acq_date"][5:7]) == m]; c = Counter(r["zona"] for r in rm)
            tt.append([BULAN[m - 1], str(c["Hutan Desa"]), str(c["Desa Nanga Lauk"]), str(c["Radius 5 km"]), str(c["Radius 5-25 km"]),
                       fmt(sum(c[z] for z in ZONES[:5]))])
    # per petak
    pc = Counter(r["petak"] for r in near if r["petak"])
    pt = [["Petak", "Deteksi", "Tanggal deteksi (WIB)"]] + [[p, str(n), ", ".join(sorted({r["waktu_wib"][:10] for r in near if r["petak"] == p}))]
                                                       for p, n in sorted(pc.items(), key=lambda x: -x[1])]
    detail = [["No", "Waktu (WIB)", "Sensor", "Lat", "Lon", "Keyakinan", "FRP (MW)", "Zona", "Jarak (km)", "Petak", "Kawasan", "Tutupan 2026", "Gambut", "Status"]]
    for i, r in enumerate(near, 1):
        detail.append([str(i), r["waktu_wib"], fam(r), f'{num(r["latitude"]):.5f}', f'{num(r["longitude"]):.5f}', conf_txt(r),
                       f'{num(r["frp"]):.1f}', r["zona"], f'{num(r["jarak_desa_km"]):.2f}', r["petak"], (r["fungsi_kawasan"] or "").split(" (")[0],
                       r["tutupan_2026"], (r["gambut"] or "").split(" (")[0], r["status_data"]])
    ringt = None
    if year is not None and ring:
        ringt = [["No", "Waktu (WIB)", "Sensor", "Lat", "Lon", "Keyakinan", "FRP (MW)", "Jarak (km)", "Status"]]
        for i, r in enumerate(ring, 1):
            ringt.append([str(i), r["waktu_wib"], fam(r), f'{num(r["latitude"]):.5f}', f'{num(r["longitude"]):.5f}', conf_txt(r),
                          f'{num(r["frp"]):.1f}', f'{num(r["jarak_desa_km"]):.1f}', r["status_data"]])
    et = [["ID kejadian", "Mulai (WIB)", "Selesai (WIB)", "Deteksi", "Zona", "Petak", "FRP maks", "Cakupan piksel (ha)"]]
    for e in sorted(evs, key=lambda e: e["mulai_wib"]):
        et.append([e["id_kejadian"], e["mulai_wib"], e["selesai_wib"], str(e["jumlah_deteksi"]), e["zona_terdekat"], e["petak"],
                   f'{e["frp_maks_mw"]:.1f}', f'{e["luas_cakupan_piksel_ha"]:.1f}'])
    n_sp = sum(1 for r in rows if r["status_data"] == "SP")
    # ringkasan
    title = f"Laporan Titik Panas Desa & Hutan Desa Nanga Lauk {year}" if year else "Laporan Titik Panas Desa & Hutan Desa Nanga Lauk 2001–Sekarang"
    peak = Counter(r["acq_date"][:4] for r in focus).most_common(1)
    peak_m = Counter(int(r["acq_date"][5:7]) for r in focus).most_common(1)
    hd_list = [r for r in near if r["zona"] == "Hutan Desa"]
    summ = [
        f"Periode data: {span[0]} s.d. {span[1]}; {fmt(len(rows))} deteksi titik panas di kotak Kapuas Hulu dan sekitarnya "
        f"({fmt(n_sp)} dari arsip final NASA/SP, {fmt(len(rows) - n_sp)} Near Real-Time).",
        f"Di dalam Hutan Desa: {zc['Hutan Desa']} deteksi" + (" (" + "; ".join(f'{r["waktu_wib"]} WIB, {fam(r)}, petak {r["petak"] or "-"}' for r in hd_list) + ")." if hd_list else "."),
        f"Di dalam Desa Nanga Lauk (di luar Hutan Desa): {zc['Desa Nanga Lauk']} deteksi; radius 0–5 km dari batas desa: {zc['Radius 5 km']}; radius 5–25 km: {fmt(zc['Radius 5-25 km'])}.",
        f"Kejadian api (pengelompokan deteksi ≤ 1 km dan ≤ 2 hari) di desa dan radius 5 km: {len(evs)} kejadian.",
    ]
    if peak:
        summ.append((f"Tahun dengan deteksi terbanyak dalam radius 25 km: {peak[0][0]} ({peak[0][1]} deteksi). " if year is None else "")
                    + (f"Bulan dengan deteksi terbanyak: {BULAN[peak_m[0][0] - 1]}." if peak_m else ""))
    return dict(title=title, period=lbl, span=span, summary=summ, figs=F, zone_table=zt, time_table=tt, petak_table=pt,
                detail=detail, ring=ringt, events=et, year=year)

METHOD = [
    ("Sumber data", "Deteksi titik panas (active fire) dari NASA Fire Information for Resource Management System (FIRMS): VIIRS 375 m "
     "(S-NPP sejak 2012, NOAA-20 sejak 2018, NOAA-21) dan MODIS 1 km (Terra/Aqua sejak 2000). Data 2001–2025 berasal dari arsip tahunan final "
     "(Standard Processing, SP); data terbaru dari Near Real-Time (NRT). Saat arsip final terbit, data NRT otomatis diganti."),
    ("Pengolahan", "Nilai asli FIRMS tidak diubah. Setiap deteksi ditambahi: zona dan jarak ke batas Desa Nanga Lauk (dihitung pada WGS 1984 "
     "UTM Zona 49N), kode petak grid patroli 250 m, kabupaten, fungsi kawasan hutan, tutupan lahan 2026, dan kedalaman gambut. Deteksi bertipe "
     "gunung api, sumber panas statis (industri), dan lepas pantai (type 1–3) tidak dihitung."),
    ("Zona", "Hutan Desa (SK.685/Menlhk-PSKL/PKPS/PSL.0/2/2017); Desa Nanga Lauk (batas desa BIG 2020); radius 5 km dan 5–25 km dari batas desa; "
     "Kapuas Hulu (batas kabupaten BIG 2021)."),
    ("Kejadian api", "Deteksi yang berjarak ≤ 1 km dan ≤ 2 hari dikelompokkan menjadi satu kejadian. Luas cakupan piksel adalah gabungan "
     "ukuran piksel satelit (scan × track), BUKAN luas terbakar."),
    ("Keterbatasan", "Titik panas adalah anomali panas yang terdeteksi satelit saat melintas (beberapa kali per hari), bukan konfirmasi kebakaran; "
     "api kecil atau di bawah tajuk dapat tidak terdeteksi, dan awan menghalangi deteksi. Jumlah deteksi antar-sensor tidak sebanding "
     "(VIIRS lebih peka dari MODIS), sehingga tren dibandingkan dalam sensor yang sama. Luas terbakar diukur dengan analisa citra (burn scar) "
     "atau pengukuran lapangan."),
    ("Sitasi", "We acknowledge the use of data and imagery from NASA's Fire Information for Resource Management System (FIRMS) "
     "(https://earthdata.nasa.gov/firms), part of NASA's Earth Science Data and Information System (ESDIS). DOI: VIIRS S-NPP "
     "10.5067/FIRMS/VIIRS/VNP14IMGT_NRT.002; VIIRS NOAA-20 10.5067/FIRMS/VIIRS/VJ114IMGT_NRT.002; VIIRS NOAA-21 10.5067/VIIRS/VJ214IMG_NRT.002; "
     "MODIS 10.5067/FIRMS/MODIS/MCD14DL.NRT.0061."),
]
DICT = [["Kolom", "Keterangan"], ["Waktu (WIB)", "Waktu lintasan satelit (acq_date + acq_time UTC + 7 jam)"],
        ["Sensor", "MODIS, VIIRS S-NPP, VIIRS NOAA-20, VIIRS NOAA-21"], ["Lat, Lon", "Pusat piksel deteksi, WGS 1984 (derajat desimal)"],
        ["Keyakinan", "VIIRS: rendah/nominal/tinggi; MODIS: 0–100%"], ["FRP (MW)", "Fire Radiative Power, kekuatan radiasi panas"],
        ["Zona, Jarak", "Posisi terhadap Hutan Desa/desa dan jarak ke batas desa (km)"], ["Petak", "Kode grid patroli 250 m Hutan Desa"],
        ["Kawasan, Tutupan, Gambut", "Fungsi kawasan hutan, tutupan lahan 2026, kedalaman gambut di lokasi deteksi"],
        ["Status", "SP = arsip final NASA; NRT = Near Real-Time (sementara)"]]

# ------------------------------------------------------------------ Word
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def _table(doc, data, widths=None, size=8, header_repeat=True):
    t = doc.add_table(rows=len(data), cols=len(data[0])); t.style = "Table Grid"
    for i, (row, trow) in enumerate(zip(data, t.rows)):
        cells = trow.cells
        for j, v in enumerate(row):
            c = cells[j]
            p = c.paragraphs[0]; run = p.add_run(str(v)); run.font.size = Pt(size)
            if i == 0:
                run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); _shade(c, HEAD.lstrip("#"))
            elif i % 2 == 0:
                _shade(c, "F1F4F1")
    if header_repeat:
        trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    return t

def build_docx(C, path, audit):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in ("Heading 1", "Heading 2", "Title"):
        doc.styles[s].font.name = "Arial"; doc.styles[s].font.color.rgb = RGBColor(0x22, 0x38, 0x2B)
    sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
    doc.add_heading(C["title"], 0)
    p = doc.add_paragraph(f'Periode: {C["span"][0]} s.d. {C["span"][1]}  ·  Dibuat otomatis: {datetime.datetime.now(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC  ·  '
                          f'Sumber: NASA FIRMS'); p.runs[0].font.size = Pt(9); p.runs[0].font.color.rgb = RGBColor(0x5B, 0x6A, 0x61)
    doc.add_heading("1. Ringkasan", 1)
    for s in C["summary"]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("2. Sumber data dan metode", 1)
    for k, v in METHOD:
        p = doc.add_paragraph(); r = p.add_run(k + ". "); r.bold = True; p.add_run(v)
    doc.add_heading("3. Peta", 1)
    doc.add_picture(C["figs"]["map25"], width=Cm(16)); doc.add_picture(C["figs"]["mapdesa"], width=Cm(16))
    doc.add_heading("4. Grafik", 1)
    for k in ("annual_modis", "annual_viirs", "monthly", "hour", "context"):
        if k in C["figs"]:
            doc.add_picture(C["figs"][k], width=Cm(16.5))
    doc.add_heading("5. Rekapitulasi", 1)
    doc.add_paragraph("Tabel 1. Jumlah deteksi per zona dan sensor")
    _table(doc, C["zone_table"])
    doc.add_paragraph("")
    doc.add_paragraph("Tabel 2. Jumlah deteksi per " + ("tahun" if C["year"] is None else "bulan") + " (*Kapuas Hulu = semua zona di Kapuas Hulu)")
    _table(doc, C["time_table"])
    if len(C["petak_table"]) > 1:
        doc.add_paragraph(""); doc.add_paragraph("Tabel 3. Petak patroli yang pernah terdeteksi titik panas")
        _table(doc, C["petak_table"])
    doc.add_heading("6. Kejadian api di desa dan radius 5 km", 1)
    if len(C["events"]) > 1:
        _table(doc, C["events"], size=7.5)
    else:
        doc.add_paragraph("Tidak ada kejadian pada periode ini.")
    # lampiran landscape
    new = doc.add_section(); new.orientation = WD_ORIENT.LANDSCAPE
    new.page_width, new.page_height = sec.page_height, sec.page_width
    new.left_margin = new.right_margin = Cm(1.5)
    doc.add_heading("Lampiran A. Daftar detail deteksi di Hutan Desa, Desa, dan radius 5 km", 1)
    if len(C["detail"]) > 1:
        _table(doc, C["detail"], size=7)
    else:
        doc.add_paragraph("Tidak ada deteksi pada periode ini.")
    if C["ring"]:
        doc.add_heading("Lampiran B. Daftar deteksi radius 5–25 km", 1); _table(doc, C["ring"], size=7)
    doc.add_heading("Lampiran C. Kamus data", 1); _table(doc, DICT, size=8)
    doc.add_heading("Lampiran D. Jejak audit sumber data", 1)
    doc.add_paragraph(audit)
    doc.save(path)

# ------------------------------------------------------------------ PDF
def build_pdf(C, path, audit):
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor(INK))
    h0 = ParagraphStyle("h0", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21, spaceAfter=6, textColor=colors.HexColor(HEAD))
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor(HEAD))
    small = ParagraphStyle("s", parent=base, fontSize=8.5, textColor=colors.HexColor(MUTED))
    cellst = ParagraphStyle("c", parent=base, fontSize=7, leading=8.5)
    cellh = ParagraphStyle("ch", parent=cellst, fontName="Helvetica-Bold", textColor=colors.white)
    def esc(s): return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    def tbl(data, widths=None, fs=7):
        cs = ParagraphStyle("cx", parent=cellst, fontSize=fs, leading=fs + 1.5); ch = ParagraphStyle("chx", parent=cs, fontName="Helvetica-Bold", textColor=colors.white)
        rows = [[Paragraph(esc(v), ch if i == 0 else cs) for v in r] for i, r in enumerate(data)]
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(HEAD)), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(LINE)),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F4F1")]),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        return t
    def footer(canvas, doc_):
        canvas.saveState(); canvas.setFont("Helvetica", 7.5); canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(1.8 * cm, 1.1 * cm, C["title"]); canvas.drawRightString(doc_.pagesize[0] - 1.8 * cm, 1.1 * cm, f"Halaman {doc_.page}")
        canvas.restoreState()
    s = []
    s.append(Paragraph(esc(C["title"]), h0))
    s.append(Paragraph(f'Periode: {C["span"][0]} s.d. {C["span"][1]} · Dibuat otomatis: {datetime.datetime.now(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC · Sumber: NASA FIRMS', small))
    s.append(Paragraph("1. Ringkasan", h1))
    for t in C["summary"]:
        s.append(Paragraph("• " + esc(t), base))
    s.append(Paragraph("2. Sumber data dan metode", h1))
    for k, v in METHOD:
        s.append(Paragraph(f"<b>{esc(k)}.</b> {esc(v)}", base)); s.append(Spacer(1, 3))
    s.append(PageBreak()); s.append(Paragraph("3. Peta", h1))
    s.append(RLImage(C["figs"]["map25"], width=15.5 * cm, height=15.5 * cm * 6.4 / 7.2))
    s.append(RLImage(C["figs"]["mapdesa"], width=15.5 * cm, height=15.5 * cm * 6.4 / 7.2))
    s.append(Paragraph("4. Grafik", h1))
    dims = {"annual_modis": 3.0, "annual_viirs": 3.0, "monthly": 2.8, "hour": 2.4, "context": 2.9}
    for k, hh in dims.items():
        if k in C["figs"]:
            w = 17 * cm; s.append(RLImage(C["figs"][k], width=w, height=w * hh / (7.4 if k == "context" else 7.2))); s.append(Spacer(1, 4))
    s.append(Paragraph("5. Rekapitulasi", h1))
    s.append(Paragraph("Tabel 1. Jumlah deteksi per zona dan sensor", small)); s.append(tbl(C["zone_table"], fs=8))
    s.append(Spacer(1, 8)); s.append(Paragraph("Tabel 2. Jumlah deteksi per " + ("tahun" if C["year"] is None else "bulan") + " (*Kapuas Hulu = semua zona di Kapuas Hulu)", small))
    s.append(tbl(C["time_table"], fs=7.5))
    if len(C["petak_table"]) > 1:
        s.append(Spacer(1, 8)); s.append(Paragraph("Tabel 3. Petak patroli yang pernah terdeteksi titik panas", small))
        s.append(tbl(C["petak_table"], widths=[2 * cm, 2 * cm, 13 * cm], fs=7.5))
    s.append(Paragraph("6. Kejadian api di desa dan radius 5 km", h1))
    s.append(tbl(C["events"], fs=7) if len(C["events"]) > 1 else Paragraph("Tidak ada kejadian pada periode ini.", base))
    s.append(PageBreak()); s.append(Paragraph("Lampiran A. Daftar detail deteksi di Hutan Desa, Desa, dan radius 5 km", h1))
    if len(C["detail"]) > 1:
        w = [0.8, 2.3, 1.6, 1.4, 1.5, 1.3, 1.1, 1.8, 1.1, 1.0, 1.0, 2.0, 1.4, 0.8]
        s.append(tbl(C["detail"], widths=[x * cm for x in w], fs=6))
    else:
        s.append(Paragraph("Tidak ada deteksi pada periode ini.", base))
    if C["ring"]:
        s.append(Paragraph("Lampiran B. Daftar deteksi radius 5–25 km", h1)); s.append(tbl(C["ring"], fs=6.5))
    s.append(Paragraph("Lampiran C. Kamus data", h1)); s.append(tbl(DICT, widths=[4.5 * cm, 12.5 * cm], fs=7.5))
    s.append(Paragraph("Lampiran D. Jejak audit sumber data", h1)); s.append(Paragraph(esc(audit), small))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                      title=C["title"], author="Pantau Api Nanga Lauk").build(s, onFirstPage=footer, onLaterPages=footer)

# ------------------------------------------------------------------ utama
def audit_text():
    idx_f = os.path.join(DATA, "provenance", "arsip_tahunan_index.json")
    n = len(json.load(open(idx_f))) if os.path.exists(idx_f) else 0
    return (f"Setiap file sumber NASA FIRMS yang diunduh dicatat URL, waktu unduh (UTC), ukuran, dan sidik SHA-256-nya di folder "
            f"data/provenance pada repositori publik https://github.com/erikprcfi/titik-api-desa-nanga-lauk ({n} file arsip tahunan tercatat; "
            f"pembaruan Near Real-Time tiap 3 jam tercatat di file jsonl bulanan). Riwayat perubahan seluruh data tersimpan di riwayat commit "
            f"repositori tersebut. Setiap titik dapat dicek silang di https://firms.modaps.eosdis.nasa.gov/map dengan tanggal dan koordinat yang sama.")

def build_reports(recs, events, years, full=True):
    os.makedirs(OUT, exist_ok=True)
    audit = audit_text(); made = []
    jobs = ([None] if full else []) + sorted(years)
    for y in jobs:
        C = content(recs, events, y)
        base = "laporan_titik_panas_nanga_lauk_" + ("lengkap" if y is None else str(y))
        build_docx(C, os.path.join(OUT, base + ".docx"), audit)
        build_pdf(C, os.path.join(OUT, base + ".pdf"), audit)
        made.append(base)
    return made


if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    argv = sys.argv[:]; sys.argv = [sys.argv[0]]
    import pipeline as pl
    recs = pl.load_archive(); events = pl.cluster(recs)
    years = sorted({r["acq_date"][:4] for r in recs.values()})
    if "--tahun" in argv:
        years = [argv[argv.index("--tahun") + 1]]
    made = build_reports(recs, events, years, full=True)
    json.dump({"dibuat_utc": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "file": made},
              open(os.path.join(OUT, "daftar.json"), "w"), indent=1)
    print(len(made), "laporan dibuat")
